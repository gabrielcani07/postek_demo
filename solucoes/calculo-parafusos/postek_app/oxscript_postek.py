from ox_script import *
import ctypes
from decimal import Decimal, InvalidOperation, ROUND_FLOOR, ROUND_HALF_UP
import fcntl
import os
import re
import time


EXCEL_PATH = "produtos_parafusos.xlsx"
LAYOUT_PATH = "layout_parafusos.prn"

produtos = []
produto_atual = None
peso_total_g = Decimal("0")
quantidade = 0
sobra_g = Decimal("0")
valor_total = Decimal("0")
status_balanca = "Aguardando peso pela USB da balanca."
ultimo_peso_lido_g = None
porta_balanca_nome = None
dado_bruto_balanca = "-"
ch340_fd = None
ch340_endpoints = []
ch340_buffer = ""
ch340_pacotes_processados = 0
DEBUG_BALANCA = False
TEMPO_CAPTURA_PACOTE_S = 2.5
TEMPO_MAX_CAPTURA_PACOTE_S = 5.0
PESO_MAXIMO_VALIDO_KG = Decimal("32")
PORTAS_BALANCA = (
    "PORT_SERIAL",
    "PORT_USBHOST",
    "PORT_USBDEVICE",
    "PORT_RS232",
    "PORT_COM",
    "PORT_COM1",
    "PORT_UART",
)

CH340_VID = "1a86"
CH340_PIDS = ("7523", "5523", "7522", "e523")
USB_DIR_IN = 0x80
USB_TYPE_VENDOR = 0x40
USB_RECIP_DEVICE = 0x00
CH341_REQ_READ_VERSION = 0x5F
CH341_REQ_WRITE_REG = 0x9A
CH341_REQ_SERIAL_INIT = 0xA1
CH341_REQ_MODEM_CTRL = 0xA4
CH341_REG_PRESCALER = 0x12
CH341_REG_DIVISOR = 0x13
CH341_REG_LCR = 0x18
CH341_REG_LCR2 = 0x25
CH341_LCR_ENABLE_RX = 0x80
CH341_LCR_ENABLE_TX = 0x40
CH341_LCR_CS8 = 0x03
CH341_BIT_RTS = 1 << 6
CH341_BIT_DTR = 1 << 5


class CtrlTransfer(ctypes.Structure):
    _fields_ = [
        ("bRequestType", ctypes.c_uint8),
        ("bRequest", ctypes.c_uint8),
        ("wValue", ctypes.c_uint16),
        ("wIndex", ctypes.c_uint16),
        ("wLength", ctypes.c_uint16),
        ("timeout", ctypes.c_uint32),
        ("data", ctypes.c_void_p),
    ]


class BulkTransfer(ctypes.Structure):
    _fields_ = [
        ("ep", ctypes.c_uint),
        ("len", ctypes.c_uint),
        ("timeout", ctypes.c_uint),
        ("data", ctypes.c_void_p),
    ]


def ioctl_number(direction, number, size):
    return (direction << 30) | (ord("U") << 8) | number | (size << 16)


USBDEVFS_CONTROL = ioctl_number(3, 0, ctypes.sizeof(CtrlTransfer))
USBDEVFS_BULK = ioctl_number(3, 2, ctypes.sizeof(BulkTransfer))
USBDEVFS_CLAIMINTERFACE = ioctl_number(2, 15, ctypes.sizeof(ctypes.c_uint))


def log_debug(texto):
    if DEBUG_BALANCA:
        print("[BALANCA] " + str(texto))


def decimal_ptbr(valor):
    texto = str(valor).strip().replace(",", ".")
    if texto == "":
        return Decimal("0")
    return Decimal(texto)


def moeda_ptbr(valor):
    return ("R$ " + format(valor, ".2f")).replace(".", ",")


def quantidade_formatada(valor):
    return str(valor) + " uni"


def peso_formatado(valor_g):
    if valor_g < Decimal("1000"):
        valor = valor_g.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)
        texto = format(valor, "f").rstrip("0").rstrip(".")
        return texto.replace(".", ",") + " g"

    valor_kg = (valor_g / Decimal("1000")).quantize(
        Decimal("0.001"), rounding=ROUND_HALF_UP
    )
    return format(valor_kg, ".3f").replace(".", ",") + " kg"


def limpar_texto_balanca(texto):
    texto_limpo = ""
    for caractere in str(texto):
        if caractere == "\n" or caractere == "\r" or caractere == "\t":
            texto_limpo += " "
        elif ord(caractere) >= 32:
            texto_limpo += caractere
    return " ".join(texto_limpo.split())


def pacotes_peso_urano(texto):
    texto_normalizado = str(texto).lower().replace(",", ".")
    pesos = re.findall(
        r"peso\s*l\s*:\s*([-+]?\d+(?:\.\d+)?)\s*kg", texto_normalizado
    )
    if len(pesos) > 0:
        return pesos

    pesos = re.findall(r"t2bn0\s+([-+]?\d+(?:\.\d+)?)\s*kg", texto_normalizado)
    if len(pesos) > 0:
        return pesos

    texto_sem_tara = re.sub(
        r"tara\s*:\s*[-+]?\d+(?:\.\d+)?\s*kg", " ", texto_normalizado
    )
    pesos = re.findall(r"([-+]?\d+\.\d+)\s*kg", texto_sem_tara)
    pesos = [
        peso
        for peso in pesos
        if Decimal(peso.replace("+", "")) > Decimal("0")
        and Decimal(peso.replace("+", "")) <= PESO_MAXIMO_VALIDO_KG
    ]
    if len(pesos) > 0:
        return pesos

    texto_compacto = re.sub(r"\s+", "", texto_normalizado)
    if re.fullmatch(r"(?:\d{5})+", texto_compacto):
        grupos_numericos = re.findall(r"\d{5}", texto_compacto)
        return [str(Decimal(grupo) / Decimal("1000")) for grupo in grupos_numericos]

    return []


def extrair_peso_urano_g(dado):
    if isinstance(dado, bytes):
        dado = dado.decode("utf-8", errors="ignore")

    texto = str(dado).strip()
    if texto == "":
        raise ValueError("Balanca nao retornou peso.")

    texto_normalizado = texto.lower().replace(",", ".")

    pesos_liquidos = pacotes_peso_urano(texto_normalizado)
    if len(pesos_liquidos) > 0:
        try:
            peso_kg = Decimal(pesos_liquidos[-1].replace("+", ""))
        except InvalidOperation as erro:
            raise ValueError("Peso liquido invalido recebido da balanca: " + texto) from erro
        return (peso_kg * Decimal("1000")).quantize(Decimal("0.001"))

    peso_com_unidade = re.search(r"([-+]?\d+(?:\.\d+)?)\s*kg", texto_normalizado)
    if peso_com_unidade is not None:
        try:
            peso_kg = Decimal(peso_com_unidade.group(1).replace("+", ""))
        except InvalidOperation as erro:
            raise ValueError("Peso invalido recebido da balanca: " + texto) from erro
        return (peso_kg * Decimal("1000")).quantize(Decimal("0.001"))

    peso_com_unidade = re.search(r"([-+]?\d+(?:\.\d+)?)\s*g", texto_normalizado)
    if peso_com_unidade is not None:
        try:
            peso_g = Decimal(peso_com_unidade.group(1).replace("+", ""))
        except InvalidOperation as erro:
            raise ValueError("Peso invalido recebido da balanca: " + texto) from erro
        return peso_g.quantize(Decimal("0.001"))

    matches = re.findall(r"[-+]?\d+(?:\.\d+)?", texto_normalizado)
    if len(matches) == 0:
        raise ValueError("Nao foi possivel identificar peso em: " + texto)

    try:
        peso = Decimal(matches[-1].replace("+", ""))
    except InvalidOperation as erro:
        raise ValueError("Peso invalido recebido da balanca: " + texto) from erro

    return peso.quantize(Decimal("0.001"))


def ler_arquivo(caminho):
    try:
        with open(caminho, "r") as arquivo:
            return arquivo.read().strip()
    except Exception:
        return ""


def encontrar_ch340():
    base = "/sys/bus/usb/devices"
    if not os.path.exists(base):
        return None

    for nome in os.listdir(base):
        caminho = os.path.join(base, nome)
        vid = ler_arquivo(os.path.join(caminho, "idVendor")).lower()
        pid = ler_arquivo(os.path.join(caminho, "idProduct")).lower()
        if vid == CH340_VID and pid in CH340_PIDS:
            log_debug("CH340 encontrado em " + caminho + " VID:PID=" + vid + ":" + pid)
            return caminho
    log_debug("CH340 nao encontrado em /sys/bus/usb/devices")
    return None


def caminho_usbfs(caminho_sysfs):
    busnum = int(ler_arquivo(os.path.join(caminho_sysfs, "busnum")))
    devnum = int(ler_arquivo(os.path.join(caminho_sysfs, "devnum")))
    return "/dev/bus/usb/%03d/%03d" % (busnum, devnum)


def endpoints_bulk_in(caminho_sysfs):
    endpoints = []
    base_nome = os.path.basename(caminho_sysfs)
    pasta_base = os.path.dirname(caminho_sysfs)
    for nome in os.listdir(pasta_base):
        if not nome.startswith(base_nome + ":"):
            continue
        interface = os.path.join(pasta_base, nome)
        for raiz, dirs, arquivos in os.walk(interface):
            if "bEndpointAddress" not in arquivos:
                continue
            endereco = int(ler_arquivo(os.path.join(raiz, "bEndpointAddress")), 16)
            atributos = int(ler_arquivo(os.path.join(raiz, "bmAttributes")), 16)
            if endereco & 0x80 and (atributos & 0x03) == 0x02:
                endpoints.append(endereco)
    if len(endpoints) == 0:
        endpoints.append(0x82)
    return endpoints


def control_out(fd, request, value, index):
    transfer = CtrlTransfer(
        USB_TYPE_VENDOR | USB_RECIP_DEVICE,
        request,
        value,
        index,
        0,
        1000,
        None,
    )
    return fcntl.ioctl(fd, USBDEVFS_CONTROL, transfer)


def control_in(fd, request, value, index, tamanho):
    buffer = ctypes.create_string_buffer(tamanho)
    transfer = CtrlTransfer(
        USB_DIR_IN | USB_TYPE_VENDOR | USB_RECIP_DEVICE,
        request,
        value,
        index,
        tamanho,
        1000,
        ctypes.cast(buffer, ctypes.c_void_p),
    )
    fcntl.ioctl(fd, USBDEVFS_CONTROL, transfer)
    return buffer.raw


def bulk_in(fd, endpoint, tamanho=4096, timeout=100):
    buffer = ctypes.create_string_buffer(tamanho)
    transfer = BulkTransfer(
        endpoint,
        tamanho,
        timeout,
        ctypes.cast(buffer, ctypes.c_void_p),
    )
    lido = fcntl.ioctl(fd, USBDEVFS_BULK, transfer)
    return buffer.raw[:lido]


def claim_interface(fd, numero=0):
    interface = ctypes.c_uint(numero)
    return fcntl.ioctl(fd, USBDEVFS_CLAIMINTERFACE, interface)


def inicializar_ch340(fd):
    versao = control_in(fd, CH341_REQ_READ_VERSION, 0, 0, 2)
    version = versao[0]
    control_out(fd, CH341_REQ_SERIAL_INIT, 0, 0)
    control_out(
        fd,
        CH341_REQ_WRITE_REG,
        (CH341_REG_DIVISOR << 8) | CH341_REG_PRESCALER,
        0xB282,
    )
    if version >= 0x30:
        control_out(
            fd,
            CH341_REQ_WRITE_REG,
            (CH341_REG_LCR2 << 8) | CH341_REG_LCR,
            CH341_LCR_ENABLE_RX | CH341_LCR_ENABLE_TX | CH341_LCR_CS8,
        )
    control_out(fd, CH341_REQ_MODEM_CTRL, (~(CH341_BIT_RTS | CH341_BIT_DTR)) & 0xFFFF, 0)
    return version


def abrir_ch340():
    global ch340_fd
    global ch340_endpoints

    if ch340_fd is not None:
        return True

    caminho_sysfs = encontrar_ch340()
    if caminho_sysfs is None:
        return False

    caminho_usb = caminho_usbfs(caminho_sysfs)
    log_debug("Abrindo USB raw " + caminho_usb)
    fd = os.open(caminho_usb, os.O_RDWR)
    claim_interface(fd, 0)
    versao = inicializar_ch340(fd)
    ch340_fd = fd
    ch340_endpoints = endpoints_bulk_in(caminho_sysfs)
    log_debug("CH340 inicializado versao=0x%02x endpoints=%s" % (versao, ch340_endpoints))
    limpar_buffer_ch340()
    return True


def limpar_buffer_ch340():
    if ch340_fd is None:
        return

    fim = time.time() + 0.5
    while time.time() < fim:
        teve_dado = False
        for endpoint in ch340_endpoints:
            try:
                dados = bulk_in(ch340_fd, endpoint, timeout=50)
            except Exception:
                continue
            if dados:
                teve_dado = True
        if not teve_dado:
            break


def ler_fragmento_ch340(timeout=100):
    partes = []
    for endpoint in ch340_endpoints:
        try:
            dados = bulk_in(ch340_fd, endpoint, timeout=timeout)
        except Exception as erro:
            if getattr(erro, "errno", None) != 110:
                log_debug("bulk_in endpoint " + str(endpoint) + " erro: " + str(erro))
            continue
        if dados:
            texto = dados.decode("utf-8", errors="ignore")
            partes.append(texto)
    return "".join(partes)


def capturar_pacote_ch340(timeout=100):
    primeiro_fragmento = ler_fragmento_ch340(timeout=timeout)
    if not primeiro_fragmento:
        return ""

    pacote = primeiro_fragmento
    log_debug("Inicio pacote: " + limpar_texto_balanca(primeiro_fragmento))

    inicio = time.time()
    prazo_silencio = time.time() + TEMPO_CAPTURA_PACOTE_S
    while time.time() < prazo_silencio and time.time() - inicio < TEMPO_MAX_CAPTURA_PACOTE_S:
        fragmento = ler_fragmento_ch340(timeout=150)
        if not fragmento:
            continue

        pacote += fragmento
        prazo_silencio = time.time() + 0.6
        log_debug("Fragmento pacote: " + limpar_texto_balanca(fragmento))

        if "PESO L" in pacote and "TOTAL R$" in pacote:
            break

    pacote_limpo = limpar_texto_balanca(pacote)
    log_debug("Pacote capturado: " + pacote_limpo)
    return pacote


def ler_peso_ch340_raw(timeout=100):
    global ch340_buffer
    global ch340_pacotes_processados
    global dado_bruto_balanca

    if not abrir_ch340():
        log_debug("Nao abriu CH340")
        return None

    pacote = capturar_pacote_ch340(timeout=timeout)
    if not pacote:
        return None

    ch340_buffer = (ch340_buffer + pacote)[-2500:]
    pesos_liquidos = pacotes_peso_urano(ch340_buffer)
    log_debug("Pacotes PESO L no buffer=" + str(len(pesos_liquidos)) + " processados=" + str(ch340_pacotes_processados))
    if len(pesos_liquidos) <= ch340_pacotes_processados:
        dado_bruto_balanca = "Pacote sem PESO L completo"
        log_debug(dado_bruto_balanca)
        return None

    ch340_pacotes_processados = len(pesos_liquidos)
    dado_bruto_balanca = "Peso recebido: " + pesos_liquidos[-1] + "kg"
    log_debug(dado_bruto_balanca)
    return Decimal(pesos_liquidos[-1].replace("+", "")) * Decimal("1000")


def portas_balanca_disponiveis():
    portas = []
    for nome in PORTAS_BALANCA:
        if nome in globals():
            portas.append((nome, globals()[nome]))
    return portas


def ler_peso_balanca(timeout=100):
    global porta_balanca_nome
    global dado_bruto_balanca

    try:
        peso_ch340_g = ler_peso_ch340_raw(timeout=timeout)
        if peso_ch340_g is not None:
            log_debug("Peso via CH340_RAW_USB=" + str(peso_ch340_g))
            return "CH340_RAW_USB", peso_ch340_g
    except Exception as erro:
        dado_bruto_balanca = "Erro CH340 raw: " + str(erro)
        log_debug(dado_bruto_balanca)

    if porta_balanca_nome is not None and porta_balanca_nome in globals():
        dado = PTK_GetPortData(globals()[porta_balanca_nome], timeout=timeout)
        if dado != -1:
            dado_bruto_balanca = limpar_texto_balanca(dado)
            return porta_balanca_nome, extrair_peso_urano_g(dado)

    for nome, porta in portas_balanca_disponiveis():
        dado = PTK_GetPortData(porta, timeout=timeout)
        if dado == -1:
            continue

        porta_balanca_nome = nome
        dado_bruto_balanca = limpar_texto_balanca(dado)
        return nome, extrair_peso_urano_g(dado)

    return None, None


def carregar_produtos():
    from openpyxl import load_workbook

    workbook = load_workbook(EXCEL_PATH, data_only=True)
    sheet = workbook.active
    cabecalho = [str(celula.value).strip() for celula in sheet[1]]

    lista = []
    for linha in sheet.iter_rows(min_row=2, values_only=True):
        if not linha or linha[0] is None:
            continue
        item = dict(zip(cabecalho, linha))
        lista.append(
            {
                "sku": str(item["sku"]),
                "tipo": str(item["tipo"]),
                "codigo_barras": str(item["codigo_barras"]),
                "descricao": str(item["descricao"]),
                "medidas": str(item["medidas"]),
                "peso_unitario_g": decimal_ptbr(item["peso_unitario_g"]),
                "preco_unitario": decimal_ptbr(item["preco_unitario"]),
            }
        )
    return lista


def nome_produto(produto):
    return produto["sku"] + " - " + produto["medidas"]


def selecionar_produto(value):
    global produto_atual
    if len(produtos) == 0:
        produto_atual = None
        return
    produto_atual = produtos[int(value)]
    atualizar_resumo()


def atualizar_peso(value):
    try:
        aplicar_peso_total(decimal_ptbr(value))
    except Exception as erro:
        resumo_controller.update("Peso invalido: " + str(erro))


def aplicar_peso_total(novo_peso_g):
    global peso_total_g
    global quantidade
    global sobra_g
    global valor_total

    if produto_atual is None:
        return

    peso_total_g = novo_peso_g
    if peso_total_g <= 0:
        quantidade = 0
        sobra_g = Decimal("0")
        valor_total = Decimal("0")
        atualizar_resumo()
        return

    quantidade_decimal = peso_total_g / produto_atual["peso_unitario_g"]
    quantidade = int(quantidade_decimal.to_integral_value(rounding=ROUND_FLOOR))
    peso_calculado = Decimal(quantidade) * produto_atual["peso_unitario_g"]
    sobra_g = peso_total_g - peso_calculado
    valor_total = (Decimal(quantidade) * produto_atual["preco_unitario"]).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )
    atualizar_resumo()


def zerar_peso_resumo():
    global peso_total_g
    global quantidade
    global sobra_g
    global valor_total
    global status_balanca
    global ultimo_peso_lido_g
    global ch340_buffer
    global ch340_pacotes_processados

    peso_total_g = Decimal("0")
    quantidade = 0
    sobra_g = Decimal("0")
    valor_total = Decimal("0")
    ultimo_peso_lido_g = None
    ch340_buffer = ""
    ch340_pacotes_processados = 0
    peso_controller.update("0")
    status_balanca = "Etiqueta impressa. Aguardando proximo peso."
    log_debug("Resumo zerado apos impressao")
    atualizar_resumo()


def ler_balanca_uma_vez():
    global status_balanca
    global ultimo_peso_lido_g

    try:
        porta_nome, peso_lido_g = ler_peso_balanca()
        if peso_lido_g is None:
            status_balanca = "Sem dado novo. Verifique cabo USB-serial e protocolo."
            atualizar_resumo()
            return

        ultimo_peso_lido_g = peso_lido_g
        peso_controller.update(format(peso_lido_g, "f").rstrip("0").rstrip("."))
        status_balanca = "Peso lido em " + porta_nome + ": " + peso_formatado(
            peso_lido_g
        )
        aplicar_peso_total(peso_lido_g)
    except Exception as erro:
        status_balanca = "Nao foi possivel ler a balanca: " + str(erro)
        atualizar_resumo()


def monitorar_balanca():
    global status_balanca
    global ultimo_peso_lido_g

    while True:
        try:
            porta_nome, peso_lido_g = ler_peso_balanca(timeout=100)
            if peso_lido_g is not None and peso_lido_g != ultimo_peso_lido_g:
                log_debug("Novo peso exibido " + str(peso_lido_g) + " pela porta " + str(porta_nome))
                ultimo_peso_lido_g = peso_lido_g
                peso_controller.update(format(peso_lido_g, "f").rstrip("0").rstrip("."))
                status_balanca = peso_formatado(peso_lido_g)
                aplicar_peso_total(peso_lido_g)
            elif peso_lido_g is not None:
                log_debug("Peso repetido ignorado: " + str(peso_lido_g))
        except Exception as erro:
            status_balanca = "Erro na leitura USB: " + str(erro)
            log_debug("Erro monitorar_balanca: " + str(erro))
            atualizar_resumo()
            time.sleep(1)


def atualizar_resumo():
    if produto_atual is None:
        resumo_controller.update("Nenhum produto carregado")
        return

    texto = (
        "SKU: " + produto_atual["sku"] + "\n"
        "Produto: " + produto_atual["descricao"] + "\n"
        "Medidas: " + produto_atual["medidas"] + "\n"
        "Peso unit.: " + str(produto_atual["peso_unitario_g"]) + " g\n"
        "Peso total: " + peso_formatado(peso_total_g) + "\n"
        "Quantidade: " + quantidade_formatada(quantidade) + "\n"
        "Valor total: " + moeda_ptbr(valor_total) + "\n"
        "Balanca: " + status_balanca
    )
    resumo_controller.update(texto)


def imprimir_etiqueta(zerar_depois=True):
    if produto_atual is None:
        log_debug("Impressao cancelada: produto_atual None")
        return

    if quantidade <= 0:
        log_debug("Impressao cancelada: quantidade <= 0")
        return

    log_debug("Imprimindo etiqueta peso=" + str(peso_total_g) + " quantidade=" + str(quantidade))
    cmd = PTK_UpdateAllFormVariables(
        LAYOUT_PATH,
        Input1=produto_atual["sku"],
        Input2=produto_atual["codigo_barras"],
        Input3=produto_atual["descricao"],
        Input4=produto_atual["medidas"],
        Input5=quantidade_formatada(quantidade),
        Input6=peso_formatado(peso_total_g),
        Input7=moeda_ptbr(produto_atual["preco_unitario"]),
        Input8=moeda_ptbr(valor_total),
    )
    PTK_SendCmdToPrinter(cmd)
    log_debug("Comando enviado para impressora")

    if zerar_depois:
        zerar_peso_resumo()


if __name__ == "__main__":
    produtos = carregar_produtos()
    nomes_produtos = [nome_produto(produto) for produto in produtos]
    if len(produtos) > 0:
        produto_atual = produtos[0]

    controller = PTK_UIInit(
        PTK_UIPage(
            produto_controller := PTK_UIList(
                title="Produto",
                items=nomes_produtos,
                valueType="int",
                value=0,
                Onpressed=selecionar_produto,
            ),
            peso_controller := PTK_UITextBox(
                title="Peso total em gramas",
                value="0",
            ),
            resumo_controller := PTK_UITextBox(title="Resumo", value="Carregando..."),
            PTK_UIButton(title="Imprimir", Onpressed=imprimir_etiqueta),
        ),
        require_execute_confirmation=False,
    )
    atualizar_resumo()
    monitorar_balanca()
